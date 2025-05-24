import json
import os
import random
import re

import openpyxl


def remove_digits(text):
    return re.sub(r'\d+', '', text)


def remove_letters(text):
    pattern = "[a-zA-Z]"
    return re.sub(pattern, "", text)


def get_worksheet(sheet_name='', get_all=False):
    note_path = r'C:\Users\16413\Desktop\NPEE\统计\WWII\工作日志.xlsx'
    workbook: openpyxl.Workbook = openpyxl.load_workbook(filename=note_path)
    sheet_names = workbook.sheetnames
    sheet_names.sort()
    problems = []
    if os.path.exists('NoteReview.txt'):
        with open(file='NoteReview.txt', mode='r', encoding='utf-8') as f:
            lines = f.readlines()
        lines = [line.strip() for line in lines]
        for line in lines:
            if '@' in line:
                problems.append(int(line.split('@')[0]))
            else:
                problems.append(None)
        lines = [line.split('@')[-1] for line in lines]
    else:
        lines = []
    if get_all:
        sheet_names = list(workbook.sheetnames)
        worksheets = [workbook[sheet_name] for sheet_name in sheet_names]
        problems_of_sheets = []
        for sheet_name in sheet_names:
            problems_of_sheet = []
            for index, sn in enumerate(lines):
                if sn == sheet_name:
                    problems_of_sheet.append(problems[index])
            problems_of_sheets.append(problems_of_sheet)
        return worksheets, sheet_names, problems_of_sheets
    else:
        if sheet_name == '':
            if len(lines) > 0:
                print('无指定，按LRU算法指定复习笔记')
                sheet_name = general_lru(all_candidates=sheet_names, reviewed=lines)
            else:
                print('无指定，且无既存复习记录，随机挑选复习笔记')
                sheet_name = random.choice(sheet_names)

        if sheet_name in workbook.sheetnames:
            worksheet: openpyxl.Worksheet | None = workbook[sheet_name]
            print('已找到指定名称【{}】的笔记。'.format(sheet_name))
        else:
            worksheet = None
            print('未找到指定名称【{}】的笔记！'.format(sheet_name))

        problems_of_sheet = []
        for index, sn in enumerate(lines):
            if sn == sheet_name:
                problems_of_sheet.append(problems[index])
        return worksheet, sheet_name, problems_of_sheet


def read_worksheet(worksheet):
    dim = worksheet.dimensions.split(':')
    dim[0] = 'B2'
    dim = [[remove_digits(d), remove_letters(d)] for d in dim]
    all_layers_hierarchy = {}
    for i in range(ord(dim[0][0]), ord(dim[1][0]) + 1):
        col = chr(i)
        for j in range(int(dim[0][1]), int(dim[1][1]) + 1):
            index = '{}{}'.format(col, j)
            val = worksheet[index].value
            if val is not None:
                if col not in all_layers_hierarchy.keys():
                    all_layers_hierarchy[col] = {}
                all_layers_hierarchy[col][j] = [val, {}]
    return all_layers_hierarchy


def parse2dict(all_layers_hierarchy):
    for col in all_layers_hierarchy.keys():
        if col != 'B':
            current_layer = all_layers_hierarchy[col]
            upper_layer = all_layers_hierarchy[chr(ord(col) - 1)]
            current_layer_keys = list(current_layer.keys())
            current_layer_keys.sort()
            upper_layer_keys = list(upper_layer.keys())
            upper_layer_keys.sort()
            for key in current_layer_keys:
                upper_key = key
                while upper_key not in upper_layer_keys:
                    upper_key -= 1
                upper_layer[upper_key][1][key] = current_layer[key]
    top_layer = all_layers_hierarchy['B']
    top_layer_json_str = json.dumps(top_layer)
    return all_layers_hierarchy, top_layer, top_layer_json_str


def recursive_read(entry, lines=None):
    if lines is None:
        lines = []
    # entry = [val, {}]
    lines.append(entry[0])
    sub_entry_keys = list(entry[1].keys())
    sub_entry_keys.sort()
    for sub_entry_key in sub_entry_keys:
        sub_entry = entry[1][sub_entry_key]
        lines = recursive_read(entry=sub_entry, lines=lines)
    return lines


def replace_keywords_with_underscore(text):
    pattern = r'###(.*?)###'  # 匹配###之间的内容
    replaced_text = re.sub(pattern, lambda m: '__' * len(m.group(1)), text)
    return replaced_text


def extract_keywords(text):
    pattern = r'###(.*?)###'  # 匹配###之间的内容
    keywords = re.findall(pattern, text)
    return keywords


def parse_blanks(lines: list):
    lines_with_blanks = lines.copy()
    answers = []
    contexts = []
    for i, line in enumerate(lines_with_blanks):
        answers_ = extract_keywords(text=line)
        line_with_blanks = replace_keywords_with_underscore(text=line)
        lines_with_blanks[i] = line_with_blanks
        while '__' in line_with_blanks:
            line_with_blanks = line_with_blanks.replace('__', '_')
        if '_' in line_with_blanks:
            line_with_blanks = line_with_blanks.split('_')
            context = []
            for context_ in range(len(line_with_blanks) - 1):
                context.append([line_with_blanks[context_], line_with_blanks[context_ + 1]])
        else:
            context = []
        assert len(context) == len(answers_)
        contexts += context
        answers += answers_
    lines = [line.replace('###', '') for line in lines]
    return lines_with_blanks, lines, answers, contexts


def verify_input(answer, weight, index, contexts):
    my_answer = input('>>>>第{}个空<<<<\n{}【】{}\n:'.format(index + 1, contexts[index][0], contexts[index][1]))
    print('我的回答:{}'.format(my_answer))
    print('正确回答:{}'.format(answer))
    cmd = ''
    while cmd not in ['Y', 'N']:
        cmd = input('是否正确？[Y/N]:')
    score_delta = {'Y': weight, 'N': 0.0}[cmd]
    return score_delta


def general_lru(all_candidates, reviewed):
    reviewed = reviewed.copy()
    never = list(set(all_candidates).difference(set(reviewed)))
    never.sort()
    if len(never) > 0:
        next_candidate = never[0]
    else:
        # LRU Algorithm
        last_review = {}
        reviewed.reverse()
        for reviewed_candidate in reviewed:
            if reviewed_candidate is not None:
                last_review[reviewed_candidate] = reviewed.index(reviewed_candidate)
        lrr = 0
        next_candidate = None
        i_ = -1
        while next_candidate is None:
            next_candidate = reviewed[i_]
            i_ -= 1
        for reviewed_candidate in last_review.keys():
            if reviewed_candidate is not None:
                if last_review[reviewed_candidate] >= lrr:
                    lrr = last_review[reviewed_candidate]
                    next_candidate = reviewed_candidate
    return next_candidate


if __name__ == '__main__':
    weight_ = 0.5
    while True:
        score_total = 0
        worksheet_ = None
        sheet_name_ = ''
        pos_ = []
        while worksheet_ is None:
            sheet_name_ = input('选择哪天的笔记？')
            worksheet_, sheet_name_, pos_ = get_worksheet(sheet_name=sheet_name_)
        all_layers_hierarchy_ = read_worksheet(worksheet_)
        _, top_layer_, top_layer_json_str_ = parse2dict(all_layers_hierarchy_)
        if len(pos_) <= 0:
            entry_key = random.choice(list(top_layer_.keys()))
        else:
            # 对问题也采用LRU算法
            entry_key = general_lru(all_candidates=list(top_layer_.keys()), reviewed=pos_)
        entry_ = top_layer_[entry_key]
        lines_ = recursive_read(entry=entry_)
        lines_with_blanks_, lines_, answers__, contexts_ = parse_blanks(lines=lines_)
        if len(answers__) > 0:
            lines_with_blanks_.insert(0, '=============题目=============')
            print('\n'.join(lines_with_blanks_))
            for index_, answer__ in enumerate(answers__):
                score_delta_ = verify_input(answer__, weight_, index_, contexts_)
                score_total += score_delta_
                print('目前得分:{}'.format(score_total))
        with open(file='NoteReview.txt', mode='a', encoding='utf-8') as f_:
            f_.writelines(['{}@{}\n'.format(entry_key, sheet_name_)])
