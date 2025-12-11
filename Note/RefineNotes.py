import openpyxl


def copy_sheets_simple(input_file, output_file, sheet_list):
    """简化版本，只复制数据不复制格式"""
    input_wb = openpyxl.load_workbook(input_file)
    output_wb = openpyxl.Workbook()

    # 删除默认工作表
    output_wb.remove(output_wb.active)

    for sheet_name in sheet_list:
        if sheet_name in input_wb.sheetnames:
            # 复制整个工作表
            source = input_wb[sheet_name]
            target = output_wb.create_sheet(title=sheet_name)

            # 只复制值
            for row in source.iter_rows(values_only=True):
                target.append(row)

    output_wb.save(output_file)
    print(f"已保存到: {output_file}")


if __name__ == "__main__":
    # 使用示例
    copy_sheets_simple(
        "原始文件.xlsx",
        "新文件.xlsx",
        ["Sheet1", "Sheet2"]
    )