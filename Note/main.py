import json
import re

import openpyxl
from openpyxl.utils import get_column_letter


def replace_in_excel(input_file, output_file="输出.xlsx", replace_configs=None, process_b_col=False):
    """
    在Excel中根据配置替换内容，输出到新文件，支持自动搜索定位和偏移量校准。

    参数:
        input_file: 输入Excel文件路径
        output_file: 输出Excel文件路径，默认"输出.xlsx"
        replace_configs: 替换配置列表，每个配置为字典，格式:
                        {"cell": "A1", "old": "原内容", "new": "新内容"}
    """
    if replace_configs is None:
        with open('测试.json', 'r', encoding='utf-8') as f:
            replace_configs = json.load(f)['replace']

    # 加载工作簿
    wb = openpyxl.load_workbook(input_file)
    # 处理指定的"待处理"工作表
    sheet = wb["待处理"]

    # 初始化行偏移量，用于校准后续所有单元格的行索引
    row_offset = 0

    for config in replace_configs:
        cell_pos: str = config["cell"]
        old_val = config["old"]
        new_val = config["new"]
        if not process_b_col and cell_pos.upper().startswith('B'):
            print(f"B列作为标题均不处理，跳过")
            continue
        # --- 新增：应用偏移量来修正目标单元格位置 ---
        # 1. 使用正则表达式解析单元格的列字母和行号
        match = re.match(r'^([A-Za-z]+)(\d+)$', cell_pos)
        if not match:
            print(f"警告：单元格格式 '{cell_pos}' 不正确，跳过此配置。")
            continue

        col_letter, row_num_str = match.groups()
        original_row = int(row_num_str)
        # 2. 计算校准后的行号
        corrected_row = original_row + row_offset
        corrected_cell_pos = f"{col_letter}{corrected_row}"
        # -----------------------------------------

        target_cell = None
        try:
            # 尝试获取校准后的单元格
            target_cell = sheet[corrected_cell_pos]
            current_val = target_cell.value
        except Exception as e:
            print(repr(e))
            # 如果校准后的单元格不存在，则尝试原始位置，或直接进入搜索
            print(f"信息：校准后的单元格 {corrected_cell_pos} 不存在，尝试原始位置 {cell_pos}...")
            try:
                target_cell = sheet[cell_pos]
                current_val = target_cell.value
            except Exception as e2:
                print(repr(e2))
                print(f"警告：单元格 {cell_pos} 及其校准位置均无效，将进行全局搜索。")
                current_val = None

        # 2. 检查当前单元格内容是否与old_val一致
        if current_val == old_val:
            # 内容一致，直接替换
            target_cell.value = new_val
            print(f"已在 {corrected_cell_pos if corrected_cell_pos else cell_pos} 完成替换。")
        else:
            print(
                f"单元格 {corrected_cell_pos if corrected_cell_pos else cell_pos} 内容与目标不符，开始全局搜索 '{old_val}'...")
            found = False
            max_row = sheet.max_row
            max_col = sheet.max_column

            for row in range(1, max_row + 1):
                for col in range(1, max_col + 1):
                    col_letter_found = get_column_letter(col)
                    cell_found = sheet[f"{col_letter_found}{row}"]
                    if cell_found.value == old_val:
                        # 找到匹配项，更新为新值
                        cell_found.value = new_val
                        print(f"已在 {col_letter_found}{row} 找到并替换。")

                        # --- 新增：计算并更新偏移量 ---
                        # 偏移量 = 实际找到的行号 - JSON中配置的原始行号
                        new_offset = row - original_row
                        if new_offset != row_offset:
                            print(f"发现行偏移！原始偏移: {row_offset}, 新偏移: {new_offset}。将应用新偏移量。")
                            row_offset = new_offset
                        # ---------------------------------

                        found = True
                        break
                if found:
                    break

            if not found:
                print(f"警告：未找到内容 '{old_val}'，跳过替换。")
                continue

    # 保存到新文件（不覆盖原文件）
    wb.save(output_file)
    print(f"\n所有操作完成，结果已保存至 {output_file}")


# 使用示例
if __name__ == "__main__":
    input_excel = "测试.xlsx"
    output_excel = "输出.xlsx"
    replace_in_excel(input_excel, output_excel, None, True)
